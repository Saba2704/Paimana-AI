from pptx import Presentation

pptx_path = r"C:\Users\Saba\.gemini\antigravity\brain\30f02a80-f446-4a07-9221-1a31dc33cb23\.user_uploaded\media_1790854900794.pptx"
prs = Presentation(pptx_path)

for idx, slide in enumerate(prs.slides):
    print(f"=== Slide {idx+1} ===")
    for s_idx, shape in enumerate(slide.shapes):
        txt = ""
        if shape.has_text_frame:
            txt = shape.text_frame.text.replace("\n", " ")[:60]
        print(f"  Shape {s_idx}: {shape.name} | pos: ({shape.left/914400:.2f}, {shape.top/914400:.2f}, {shape.width/914400:.2f}, {shape.height/914400:.2f}) | text: '{txt}'")
